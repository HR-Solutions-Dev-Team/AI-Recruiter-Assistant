"""
Сервис для парсинга файлов (TXT, DOCX, PDF).
"""

import io
import logging
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.core.config import settings

logger = logging.getLogger(__name__)


class FileParserError(Exception):
    """Ошибка парсинга файла."""

    pass


class FileParserService:
    """Сервис для извлечения текста из файлов."""

    SUPPORTED_EXTENSIONS = {".txt", ".docx", ".pdf"}

    def parse(self, content: bytes, filename: str) -> str:
        """
        Извлекает текст из файла.

        Args:
            content: Содержимое файла в байтах
            filename: Имя файла (для определения типа)

        Returns:
            Извлечённый текст

        Raises:
            FileParserError: Если формат не поддерживается или ошибка парсинга
        """
        ext = Path(filename).suffix.lower()

        if ext not in self.SUPPORTED_EXTENSIONS:
            raise FileParserError(
                f"Unsupported file format: {ext}. "
                f"Supported: {', '.join(self.SUPPORTED_EXTENSIONS)}"
            )

        if len(content) > settings.max_upload_size_bytes:
            raise FileParserError(
                f"File too large. Maximum size: {settings.max_upload_size_mb}MB"
            )

        try:
            if ext == ".txt":
                return self._parse_txt(content)
            elif ext == ".docx":
                return self._parse_docx(content)
            elif ext == ".pdf":
                return self._parse_pdf(content)
            else:
                raise FileParserError(f"No parser for extension: {ext}")
        except FileParserError:
            raise
        except Exception as e:
            logger.error(f"Failed to parse file {filename}: {e}")
            raise FileParserError(f"Failed to parse file: {str(e)}")

    def _parse_txt(self, content: bytes) -> str:
        """Парсит TXT файл."""
        # Пробуем разные кодировки
        for encoding in ["utf-8", "cp1251", "latin-1"]:
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue

        raise FileParserError("Failed to decode text file. Unsupported encoding.")

    def _parse_docx(self, content: bytes) -> str:
        """Парсит DOCX файл."""
        doc = Document(io.BytesIO(content))
        paragraphs = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # Также извлекаем текст из таблиц
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        return "\n".join(paragraphs)

    def _parse_pdf(self, content: bytes) -> str:
        """Парсит PDF файл."""
        reader = PdfReader(io.BytesIO(content))
        text_parts = []

        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text.strip())

        return "\n".join(text_parts)


file_parser_service = FileParserService()
